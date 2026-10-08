import json
import os
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.agents.summarizer_agent import SummarizerAgent
from app.core.config import MATERIALS_DIR, get_settings
from app.core.logger import get_logger
from app.models.db import ActionItemDB, LectureDB, MeetingDB
from app.services.rag_service import RagService

logger = get_logger("meeting_agent.materials")


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract page-by-page text from PDF slides or documents using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                pages_text.append(f"[Slide / Page {i + 1}]\n{text.strip()}")
        return "\n\n".join(pages_text)
    except Exception as e:
        logger.warning(f"Failed to extract PDF text from {file_path}: {e}")
        return ""


def extract_text_from_pptx(file_path: Path) -> str:
    """Extract slide-by-slide text from PowerPoint presentations using python-pptx."""
    try:
        from pptx import Presentation
        prs = Presentation(str(file_path))
        slides_text = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    clean = shape.text.strip()
                    if clean:
                        texts.append(clean)
            if texts:
                slides_text.append(f"[Slide {i + 1}]\n" + "\n".join(texts))
        return "\n\n".join(slides_text)
    except Exception as e:
        logger.warning(f"Failed to extract PPTX text from {file_path}: {e}")
        return ""


def extract_text_from_file(file_path: Path) -> str:
    """Detect file format and extract textual content."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    elif suffix in (".pptx", ".ppt"):
        return extract_text_from_pptx(file_path)
    elif suffix in (".md", ".txt", ".json", ".csv", ".log"):
        try:
            with open(file_path, encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception as e:
            logger.warning(f"Failed to read text file {file_path}: {e}")
            return ""
    return ""


def detect_material_type(filename: str, explicit_type: str | None = None) -> str:
    """Determine material type: presentation, summary_md, material, or note."""
    if explicit_type and explicit_type in ("presentation", "summary_md", "material", "note"):
        return explicit_type

    suffix = Path(filename).suffix.lower()
    lower_name = filename.lower()

    if suffix == ".md" or "summary" in lower_name:
        return "summary_md"
    elif suffix in (".pdf", ".pptx", ".ppt") or any(k in lower_name for k in ("folien", "slides", "presentation", "vorlesung")):
        return "presentation"
    elif suffix in (".txt", ".note"):
        return "note"
    return "material"


MAX_MATERIAL_BYTES = 50 * 1024 * 1024  # 50 MB
ALLOWED_MATERIAL_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".md", ".txt", ".docx", ".doc"}


def save_lecture_material_file(lecture_id: str, file: UploadFile, explicit_type: str | None = None) -> dict[str, Any]:
    """Save an uploaded file for a lecture with streaming chunking and size limits to prevent OOM."""
    original_name = Path(file.filename or "uploaded_material").name
    suffix = Path(original_name).suffix.lower()

    if suffix not in ALLOWED_MATERIAL_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File extension '{suffix}' not allowed. Permitted formats: {', '.join(sorted(ALLOWED_MATERIAL_EXTENSIONS))}"
        )

    target_dir = (MATERIALS_DIR / lecture_id).resolve()
    if not target_dir.is_relative_to(MATERIALS_DIR.resolve()):
        raise HTTPException(status_code=400, detail="Invalid lecture identifier")
    target_dir.mkdir(parents=True, exist_ok=True)

    unique_id = uuid.uuid4().hex[:8]
    safe_name = f"{unique_id}_{original_name}"
    saved_path = target_dir / safe_name

    total_bytes = 0
    chunk_size = 1024 * 1024  # 1MB chunks
    try:
        with open(saved_path, "wb") as f:
            while chunk := file.file.read(chunk_size):
                total_bytes += len(chunk)
                if total_bytes > MAX_MATERIAL_BYTES:
                    saved_path.unlink(missing_ok=True)
                    raise HTTPException(
                        status_code=413,
                        detail="File exceeds maximum allowed size limit of 50 MB"
                    )
                f.write(chunk)
    except HTTPException:
        saved_path.unlink(missing_ok=True)
        raise
    except Exception as e:
        saved_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"Failed to save upload: {e}") from e

    mat_type = detect_material_type(original_name, explicit_type)
    extracted_text = extract_text_from_file(saved_path)

    return {
        "id": f"mat_{uuid.uuid4().hex[:10]}",
        "title": original_name,
        "type": mat_type,
        "url": f"/api/v1/calendar/materials/file/{lecture_id}/{safe_name}",
        "filename": safe_name,
        "original_filename": original_name,
        "file_path": str(saved_path),
        "text_content": extracted_text[:80000] if extracted_text else "",
        "has_text": bool(extracted_text and len(extracted_text.strip()) > 10),
        "created_at": datetime.now(UTC).isoformat()
    }


def generate_lecture_master_summary(
    db: Session,
    lecture_id: str,
    custom_instructions: str | None = None
) -> dict[str, Any]:
    """
    Synthesize a comprehensive Master Study Summary for a lecture combining all available context:
    1. Audio / Video recording transcript
    2. Presentation slides (.pdf, .pptx)
    3. Uploaded markdown summary (.md)
    4. Student notes & additional materials
    """
    lec = db.query(LectureDB).filter(LectureDB.id == lecture_id).first()
    if not lec:
        raise HTTPException(status_code=404, detail="Lecture not found")

    context_sections = []

    # 1. Recording Transcript
    meeting = lec.meetings[0] if lec.meetings else None
    transcript_text = ""
    if meeting:
        if meeting.full_text and meeting.full_text.strip():
            transcript_text = meeting.full_text.strip()
        elif meeting.transcript_segments:
            transcript_text = "\n".join([f"{seg.speaker}: {seg.text}" for seg in meeting.transcript_segments])

    if transcript_text:
        context_sections.append(f"### [SOURCE 1: Audio / Lecture Transcript]\n{transcript_text[:35000]}")

    # 2. Materials: Presentation Slides, Uploaded MD Summaries, Reference Documents
    materials = lec.materials or []
    for mat in materials:
        mat_type = mat.get("type")
        mat_title = mat.get("title") or mat.get("original_filename") or "Document"
        text = (mat.get("text_content") or "").strip()

        # If file exists on disk but text wasn't cached, try extracting now
        if not text and mat.get("file_path") and os.path.exists(mat.get("file_path")):
            text = extract_text_from_file(Path(mat.get("file_path"))).strip()

        if text:
            if mat_type == "presentation":
                context_sections.append(f"### [SOURCE 2: Presentation Slides: {mat_title}]\n{text[:35000]}")
            elif mat_type == "summary_md":
                context_sections.append(f"### [SOURCE 3: Previous / Uploaded Markdown Summary: {mat_title}]\n{text[:30000]}")
            else:
                context_sections.append(f"### [SOURCE 4: Lecture Material / Notes: {mat_title}]\n{text[:20000]}")

    # 3. Student Notes & Description
    if lec.notes and lec.notes.strip():
        context_sections.append(f"### [SOURCE: Student Lecture Notes]\n{lec.notes.strip()}")

    if lec.description and lec.description.strip():
        context_sections.append(f"### [SOURCE: Timetable Description]\n{lec.description.strip()}")

    if not context_sections:
        raise HTTPException(
            status_code=400,
            detail="No context files, transcript, presentation slides, or notes found for this lecture to generate a summary."
        )

    combined_context = "\n\n" + ("=" * 50) + "\n\n"
    full_prompt_context = combined_context.join(context_sections)

    student_notes = (lec.notes or "").strip()

    # Detect if any official lecture recordings or presentation materials exist
    has_audio_transcript = bool(transcript_text)
    has_presentation_slides = any(mat.get("type") == "presentation" for mat in materials)
    has_uploaded_notes = any(mat.get("type") in ("summary_md", "notes") for mat in materials)
    has_official_content = has_audio_transcript or has_presentation_slides or has_uploaded_notes

    notes_system_directive = ""
    if student_notes:
        notes_system_directive = f"""
*** PRIMARY EDITORIAL DIRECTIVE: STUDENT LECTURE NOTES & CONSTRAINTS ***
The student has provided personal notes and explicit instructions for this lecture:
\"\"\"
{student_notes}
\"\"\"
CRITICAL RULES ON SCOPE & NEGATIVE CONSTRAINTS (STRICT PRIORITY):
1. NEGATIVE CONSTRAINTS OVERRIDE TEMPLATES: If the student notes or user instructions contain negative constraints (e.g. "No need to any extra info for summary!", "only focus on X", "keep brief", "no theory", "skip formulas"), you MUST strictly obey these boundaries. Do NOT include extraneous theory, formulas, or academic filler.
2. NO HALLUCINATED LECTURES: If no official lecture slides or audio recordings are provided, or if the student notes state that the lecture was skipped/unattended, DO NOT fabricate a fictitious lecture or dump unrequested textbook chapters out of thin air.
3. DIRECT PURPOSE ALIGNMENT: Focus solely on what the student specifically asked for or documented (for example: if the student noted "we need to research sources and libraries for project research", provide a concise, high-value guide for relevant sources, libraries, and next steps for that research mandate).
4. ADAPTIVE DOCUMENT STRUCTURE: Do not force standard lecture headings (like "Core Theory", "Key Formulas", or "Exam Review") if they conflict with the student's negative constraint or if there was no lecture content to extract them from. Instead, produce a clean, tailored markdown report directly matching the student's needs.
"""

    subject_name = lec.subject.name if lec.subject else "Academic Subject"
    system_prompt = f"""You are an elite academic assistant and university study synthesizer.
Your task is to synthesize a high-yield, authoritative Master Study Summary for "{lec.title}" in {subject_name}.
{notes_system_directive}

CRITICAL RULES ON SCOPE & DOMAIN FIDELITY:
1. STRICT ANTI-SOLVING DIRECTIVE: NEVER attempt to solve, implement, or write software code for student homework assignments, capstone projects, or NGO challenges mentioned in the lecture! Your role is strictly to summarize the academic knowledge, project requirements, and methodology taught — NOT to do the project work or fabricate code solutions for project briefs.
2. DOMAIN-AWARE BLUEPRINTS:
   - ONLY include programming code snippets if the lecture is an actual computer programming / software development course where code or syntax was actively taught in class.
   - For Consulting, Business, Management, Marketing, or Law lectures: DO NOT generate programming code! Instead, provide consulting frameworks, client engagement models, deliverable structures, or business matrices.

MANDATORY DOCUMENT STRUCTURE (STRICT INVERTED-PYRAMID RULE):
The "markdown_content" field MUST follow this exact, clean structure:

# {lec.title}

> **Quick Context:** 📅 {lec.start_time.strftime('%Y-%m-%d')} ({lec.start_time.strftime('%H:%M')}–{lec.end_time.strftime('%H:%M')}) | 👨‍🏫 {lec.subject.lecturer or 'Lecturer'} | 🏛️ {lec.room or 'University Room / Online'}

---

## ⚡ Executive Overview
*(CRITICAL: Keep this section SHORT and punchy — 4 to 6 lines max! No walls of text!)*
* **Core Mandate:** [1–2 sentences explaining the main objective and primary purpose of today's lecture.]
* **Key Strategic Takeaways:**
  • **Takeaway 1:** [Primary theoretical principle or core concept]
  • **Takeaway 2:** [Key trade-off, architectural rule, or consulting methodology]

---

## 📋 Tasks & Action Items for Next Lecture
*(CRITICAL: Prominent standalone checklist! Always include this section with - [ ] checkboxes outside any quotes)*
- [ ] **[Task 1 - Homework / Deliverable]:** [Specific homework, reading assignment, project milestone, or slide prep due before next lecture]
- [ ] **[Task 2 - Preparation / Setup]:** [Setup, tool installation, topic review, or group coordination for the next class]
*(If the professor did not assign explicit homework, state clear preparation recommendations for the upcoming lecture topic)*

---

## 📖 Deep Academic & Technical Foundations
*(CRITICAL: This lower section must be EXHAUSTIVE, RICH, and DEEPLY DETAILED — provide complete academic depth!)*

### 1. Pedagogical Scope & Objectives
* Why this topic was introduced and how it builds upon previous concepts.
* Core theoretical motivations and first-principles reasoning.

### 2. Architecture, Constraints & Frameworks
* Detailed rules, methodologies, platform boundaries, lifecycle stages, or industry standards discussed.

### 3. Comprehensive Concept Breakdown
* In-depth breakdown of all concepts, models, patterns, or theories presented in the lecture.
* Trade-offs, edge cases, pitfalls, and analytical comparisons.

### 4. Specifications, Definitions & Formulas
Provide a clean, structured Markdown table with formal definitions and impacts:
| Term / Metric / Concept | Formal Definition & Context | Practical / Operational Impact |
| :--- | :--- | :--- |

### 5. Practical Frameworks & Application Blueprints
- If this is a programming course: provide syntactically correct code snippets demonstrating the concepts taught in class (NEVER solutions to student project assignments).
- If this is a consulting / business course: provide the consulting methodology, matrix, engagement workflow, or deliverable template taught in class.

---

## 🎯 Exam Pointers & High-Yield Practice Questions

### ⚠️ Exam Traps & Professor's Warnings
* Specific pitfalls, exam traps, and explicit pointers on what the professor expects or will test.

### 📝 Practice Questions with Model Solutions
Provide 2–4 realistic exam-grade questions (conceptual, case analysis, or design) with complete, step-by-step model answers.
"""

    notes_block = f"*** STUDENT NOTES & DIRECTIVES (PRIORITIZE STRICTLY) ***:\n{student_notes}\n" if student_notes else ""
    custom_block = f"Additional User Instructions: {custom_instructions}\n" if custom_instructions else ""

    user_prompt = f"""Lecture Details:
Subject: {subject_name}
Title: {lec.title}
Date / Time: {lec.start_time.strftime('%Y-%m-%d %H:%M')} - {lec.end_time.strftime('%H:%M')}
{notes_block}
{custom_block}
AVAILABLE LECTURE ASSETS & CONTEXT:
{full_prompt_context}

Generate the Master Study Summary JSON. Follow the short Executive Overview at top, prominent Tasks & Action Items for Next Lecture, and rich deep academic foundations below in markdown_content."""

    settings = get_settings()
    key = settings.gemini_api_key if settings.summarization_engine == "gemini" else settings.openai_api_key
    model = settings.gemini_model if settings.summarization_engine == "gemini" else settings.openai_model
    agent = SummarizerAgent(engine=settings.summarization_engine, api_key=key, model=model)

    summary = agent.analyze(
        transcript=full_prompt_context,
        template="moodle_lecture",
        custom_instructions=user_prompt,
        system_prompt_override=system_prompt
    )

    # Unpack JSON if returned in markdown_content or executive_summary
    content_to_store = summary.markdown_content or summary.executive_summary or ""
    if content_to_store.strip().startswith("{") and content_to_store.strip().endswith("}"):
        try:
            parsed = json.loads(content_to_store.strip())
            if isinstance(parsed, dict):
                if parsed.get("markdown_content"):
                    content_to_store = parsed["markdown_content"]
                if not summary.executive_summary and parsed.get("executive_summary"):
                    summary.executive_summary = parsed["executive_summary"]
                if not summary.overview and parsed.get("overview"):
                    summary.overview = parsed["overview"]
                if not summary.key_points and parsed.get("key_points"):
                    summary.key_points = parsed["key_points"]
                if not summary.decisions and parsed.get("decisions"):
                    summary.decisions = parsed["decisions"]
                if not summary.open_questions and parsed.get("open_questions"):
                    summary.open_questions = parsed["open_questions"]
                if not summary.action_items and parsed.get("action_items"):
                    from app.models.schemas import ActionItemSchema
                    summary.action_items = [
                        ActionItemSchema(
                            task=a.get("task", ""),
                            assignee=a.get("assignee") or "Student",
                            priority=a.get("priority") or "Medium",
                            deadline=a.get("deadline"),
                            completed=False
                        )
                        for a in parsed.get("action_items", [])
                    ]
        except Exception as e:
            logger.warning(f"Failed to unpack nested JSON in master summary: {e}")

    # Store master summary into lecture override (clean markdown)
    lec.ai_summary_override = content_to_store

    # Create meeting if it doesn't exist so structured fields are preserved
    if not meeting:
        meeting = MeetingDB(
            id=f"meet_{lec.id}_{uuid.uuid4().hex[:6]}",
            lecture_id=lec.id,
            title=summary.title or lec.title,
            status="completed",
            created_at=datetime.now(UTC),
            template_used="moodle_lecture",
        )
        db.add(meeting)

    meeting.overview = summary.overview
    meeting.executive_summary = summary.executive_summary
    meeting.key_points = summary.key_points
    meeting.decisions = summary.decisions
    meeting.open_questions = summary.open_questions
    meeting.markdown_content = content_to_store

    # Refresh action items
    db.query(ActionItemDB).filter(ActionItemDB.meeting_id == meeting.id).delete()
    for ai in summary.action_items:
        db.add(ActionItemDB(
            meeting_id=meeting.id,
            task=ai.task,
            assignee=ai.assignee or "Student",
            priority=ai.priority or "Medium",
            deadline=ai.deadline,
            completed=False
        ))

    db.commit()

    # Re-index all lecture assets into RAG so grill-me and search immediately benefit
    RagService.index_lecture(db, lecture_id)

    return {
        "lecture_id": lecture_id,
        "title": summary.title or lec.title,
        "overview": summary.overview,
        "executive_summary": summary.executive_summary,
        "key_points": summary.key_points,
        "decisions": summary.decisions,
        "action_items": [ai.model_dump() for ai in summary.action_items],
        "open_questions": summary.open_questions,
        "markdown_content": content_to_store,
        "has_ai_context": True
    }
